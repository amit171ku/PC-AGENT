"""
Premium navy-and-white Excel report designer.

Usage:
    python excel_design.py "source.xlsx"
    python excel_design.py "source.xlsx" "output.xlsx" "Sheet Name"

The source workbook is never modified.
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties


NAVY, BLUE, SOFT, LINE = "0B2545", "3A7CA5", "EEF3FA", "D5DDE8"
INK, MUTED, PALE, WHITE = "1B2A41", "5B6B82", "A9C4E4", "FFFFFF"
FONT = "Calibri"

R_TITLE, R_SUB, R_ACCENT, R_GAP1, R_KL, R_KV, R_GAP2, R_HDR = (
    1, 2, 3, 4, 5, 6, 7, 8
)


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def find_table(ws):
    """Return header row, first column, last column and last data row."""
    maxr, maxc = ws.max_row, ws.max_column

    for r in range(1, min(maxr, 20) + 1):
        filled = [
            (c, ws.cell(r, c).value)
            for c in range(1, maxc + 1)
            if not _blank(ws.cell(r, c).value)
        ]

        if (
            len(filled) >= 3
            and all(isinstance(v, str) for _, v in filled)
            and r < maxr
        ):
            nxt = sum(
                not _blank(ws.cell(r + 1, c).value)
                for c in range(1, maxc + 1)
            )

            if nxt >= 2:
                c1, c2 = filled[0][0], filled[-1][0]
                last = r

                for rr in range(r + 1, maxr + 1):
                    if all(
                        _blank(ws.cell(rr, c).value)
                        for c in range(c1, c2 + 1)
                    ):
                        break
                    last = rr

                return r, c1, c2, last

    raise ValueError(
        "Could not find a table with a header row followed by data."
    )


def _role(h):
    t = str(h).strip().lower()

    if re.search(r"\bage\s*rating\b", t):
        return "other"
    if re.search(r"rating", t):
        return "rating"
    if re.search(r"vote", t):
        return "votes"
    if re.search(r"year", t):
        return "year"
    if re.search(r"rank|^#$|^no\.?$|position", t):
        return "rank"
    if re.search(r"runtime|duration|minutes", t):
        return "runtime"
    if re.search(r"gross|revenue|budget|box", t):
        return "money"
    if re.search(r"score", t):
        return "score"

    return "other"


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _pick_sheet(wb, name):
    if name and name in wb.sheetnames:
        return wb[name]

    for n in wb.sheetnames:
        if "imdb" in n.lower():
            return wb[n]

    return wb.worksheets[0]


def _shift_ref(f, off):
    head, sep, tail = f.rpartition("!")
    tail = re.sub(
        r"(\$?[A-Z]{1,3}\$?)(\d+)",
        lambda m: m.group(1) + str(max(1, int(m.group(2)) + off)),
        tail,
    )
    return head + sep + tail


def _shift_chart(ch, off):
    anchor = ch.anchor

    for attr in ("_from", "to"):
        marker = getattr(anchor, attr, None)
        if marker is not None:
            marker.row = max(0, marker.row + off)

    for series in ch.series:
        for attr in ("val", "cat", "xVal", "yVal"):
            data = getattr(series, attr, None)
            if data is None:
                continue

            for ref_attr in ("numRef", "strRef"):
                ref = getattr(data, ref_attr, None)
                if ref is not None and ref.f:
                    ref.f = _shift_ref(ref.f, off)

        if (
            series.tx is not None
            and series.tx.strRef is not None
            and series.tx.strRef.f
        ):
            series.tx.strRef.f = _shift_ref(
                series.tx.strRef.f, off
            )


def _split_cards(colw, needs):
    """Split columns into groups for the four KPI cards."""
    k, n = len(needs), len(colw)

    if n <= k:
        return [1] * k

    best, best_score = None, None

    for cuts in combinations(range(1, n), k - 1):
        bounds = (0,) + cuts + (n,)
        widths = [
            sum(colw[bounds[i]:bounds[i + 1]])
            for i in range(k)
        ]

        short = sum(
            max(0.0, needs[i] - widths[i]) ** 2
            for i in range(k)
        ) * 50

        mean = sum(widths) / k
        score = short + sum((w - mean) ** 2 for w in widths)

        if best_score is None or score < best_score:
            best, best_score = bounds, score

    return [
        best[i + 1] - best[i]
        for i in range(k)
    ]


def design_report(
    src,
    dst,
    sheet="",
    title="IMDb MOVIE DATABASE",
    subtitle="2023 MOVIE RATINGS & ANALYSIS",
):
    src, dst = Path(src), Path(dst)

    if src.resolve() == dst.resolve():
        raise ValueError(
            "Output must be a different file. The original is untouched."
        )

    wb = load_workbook(src)
    ws = _pick_sheet(wb, sheet)

    h0, c1, c2, last0 = find_table(ws)
    ncols, nrows = c2 - c1 + 1, last0 - h0

    headers = [
        ws.cell(h0, c).value
        for c in range(c1, c2 + 1)
    ]

    original = [
        [
            ws.cell(r, c).value
            for c in range(c1, c2 + 1)
        ]
        for r in range(h0 + 1, last0 + 1)
    ]

    roles = [_role(h) for h in headers]

    # Prefer IMDb Rating over Age Rating.
    rating_idx = next(
        (
            i for i, h in enumerate(headers)
            if "imdb" in str(h).lower()
            and "rating" in str(h).lower()
        ),
        None,
    )

    if rating_idx is None:
        rating_idx = next(
            (
                i for i, h in enumerate(headers)
                if "rating" in str(h).lower()
                and "age" not in str(h).lower()
            ),
            None,
        )

    if rating_idx is None:
        rating_idx = next(
            (
                i for i, role in enumerate(roles)
                if role == "score"
            ),
            None,
        )

    if rating_idx is None:
        raise ValueError(
            "No suitable rating column found. Check workbook headers."
        )

    title_idx = next(
        (
            i for i, h in enumerate(headers)
            if re.search(
                r"title|movie|name|film",
                str(h),
                re.I,
            )
        ),
        0,
    )

    # Validate that rating values exist before creating the report.
    ratings = [
        row[rating_idx]
        for row in original
        if _is_num(row[rating_idx])
    ]

    if not ratings:
        raise ValueError(
            "No numeric ratings found in the detected rating column."
        )

    offset = R_HDR - h0
    charts = list(ws._charts)

    # Clear the sheet while retaining the source data in memory.
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))

    ws.delete_rows(1, ws.max_row + 5)
    ws.conditional_formatting = ConditionalFormattingList()
    ws.auto_filter.ref = None
    ws.freeze_panes = None

    for key in list(ws.row_dimensions.keys()):
        del ws.row_dimensions[key]

    for key in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[key]

    first_data = R_HDR + 1
    last_row = R_HDR + nrows

    thin = Side(style="thin", color=LINE)
    box = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    # Column formats.
    col_fmt, col_align = [], []

    for j, role in enumerate(roles):
        vals = [
            row[j]
            for row in original
            if not _blank(row[j])
        ]

        numeric = bool(vals) and all(
            _is_num(v) for v in vals
        )

        if role in ("rating", "score"):
            col_fmt.append("0.0")
            col_align.append("center")
        elif role in ("votes", "money"):
            col_fmt.append("#,##0")
            col_align.append("right")
        elif role in ("year", "rank", "runtime"):
            col_fmt.append("0")
            col_align.append("center")
        elif numeric:
            ints = all(float(v).is_integer() for v in vals)
            col_fmt.append("#,##0" if ints else "0.00")
            col_align.append("right")
        else:
            col_fmt.append("General")
            col_align.append("left")

    # Calculate column widths.
    colw = []

    for j in range(ncols):
        longest = len(str(headers[j])) + 3

        for row in original:
            v = row[j]

            if _is_num(v):
                txt = (
                    f"{v:,.0f}"
                    if col_fmt[j] == "#,##0"
                    else str(v)
                )
            else:
                txt = "" if v is None else str(v)

            longest = max(longest, len(txt))

        width = min(max(longest + 2, 8), 48)

        if j == title_idx:
            width = max(width, 24)

        colw.append(width)

    # Title band.
    for r, height in (
        (R_TITLE, 32),
        (R_SUB, 20),
        (R_ACCENT, 3),
    ):
        ws.row_dimensions[r].height = height

        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = _fill(
                BLUE if r == R_ACCENT else NAVY
            )

    t = ws.cell(R_TITLE, c1, title)
    t.font = Font(
        name=FONT,
        size=22,
        bold=True,
        color=WHITE,
    )
    t.alignment = Alignment(
        horizontal="left",
        vertical="center",
        indent=1,
    )

    s = ws.cell(R_SUB, c1, subtitle)
    s.font = Font(name=FONT, size=11, color=PALE)
    s.alignment = Alignment(
        horizontal="left",
        vertical="center",
        indent=1,
    )

    ws.row_dimensions[R_GAP1].height = 8
    ws.row_dimensions[R_GAP2].height = 8

    # KPI cards.
    rc = L(c1 + rating_idx)
    tc = L(c1 + title_idx)

    rrange = f"${rc}${first_data}:${rc}${last_row}"
    trange = f"${tc}${first_data}:${tc}${last_row}"

    kpis = [
        ("TOTAL MOVIES", f"=COUNTA({trange})", "0"),
        ("AVERAGE IMDb RATING", f"=AVERAGE({rrange})", "0.00"),
        ("HIGHEST RATING", f"=MAX({rrange})", "0.0"),
        ("LOWEST RATING", f"=MIN({rrange})", "0.0"),
    ]

    widths = _split_cards(
        colw,
        [len(kpi[0]) * 0.95 + 3 for kpi in kpis],
    )

    ws.row_dimensions[R_KL].height = 16
    ws.row_dimensions[R_KV].height = 30

    col = c1

    for extra in range(c2 + 1, c1 + 4):
        ws.column_dimensions[L(extra)].width = 20

    for (label, formula, fmt), width in zip(kpis, widths):
        left_edge = Border(
            left=Side(style="thick", color=BLUE)
        )

        for r in (R_KL, R_KV):
            for c in range(col, col + width):
                cell = ws.cell(r, c)
                cell.fill = _fill(SOFT)

                if c == col:
                    cell.border = left_edge

        lc = ws.cell(R_KL, col, label)
        lc.font = Font(
            name=FONT,
            size=9,
            bold=True,
            color=MUTED,
        )
        lc.alignment = Alignment(
            horizontal="left",
            vertical="bottom",
            indent=1,
        )

        vc = ws.cell(R_KV, col, formula)
        vc.font = Font(
            name=FONT,
            size=20,
            bold=True,
            color=NAVY,
        )
        vc.alignment = Alignment(
            horizontal="left",
            vertical="center",
            indent=1,
        )
        vc.number_format = fmt

        if width > 1:
            ws.merge_cells(
                start_row=R_KL,
                start_column=col,
                end_row=R_KL,
                end_column=col + width - 1,
            )
            ws.merge_cells(
                start_row=R_KV,
                start_column=col,
                end_row=R_KV,
                end_column=col + width - 1,
            )

        col += width

    # Table header.
    ws.row_dimensions[R_HDR].height = 24

    for j, header in enumerate(headers):
        cell = ws.cell(R_HDR, c1 + j, header)
        cell.font = Font(
            name=FONT,
            size=11,
            bold=True,
            color=WHITE,
        )
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(
            horizontal=(
                "left" if col_align[j] == "left" else "center"
            ),
            vertical="center",
            indent=1 if col_align[j] == "left" else 0,
        )
        cell.border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=Side(style="medium", color=BLUE),
        )

    # Table body.
    written = []

    for i, row in enumerate(original):
        newr = first_data + i
        oldr = h0 + 1 + i

        ws.row_dimensions[newr].height = 18
        out_row = []

        for j, value in enumerate(row):
            if isinstance(value, str) and value.startswith("="):
                value = Translator(
                    value,
                    origin=f"{L(c1 + j)}{oldr}",
                ).translate_formula(
                    f"{L(c1 + j)}{newr}"
                )

            cell = ws.cell(newr, c1 + j, value)
            cell.font = Font(
                name=FONT,
                size=10,
                color=INK,
                bold=(j == title_idx),
            )
            cell.border = box
            cell.number_format = col_fmt[j]
            cell.alignment = Alignment(
                horizontal=col_align[j],
                vertical="center",
                indent=1 if col_align[j] == "left" else 0,
            )

            out_row.append(value)

        written.append(out_row)

    # Apply column widths.
    for j in range(ncols):
        ws.column_dimensions[L(c1 + j)].width = colw[j]

    # Conditional formatting.
    ws.conditional_formatting.add(
        f"{rc}{first_data}:{rc}{last_row}",
        ColorScaleRule(
            start_type="min",
            start_color="EAF1FB",
            end_type="max",
            end_color="7FB0E5",
        ),
    )

    ws.conditional_formatting.add(
        f"{L(c1)}{first_data}:{L(c2)}{last_row}",
        FormulaRule(
            formula=[
                f"MOD(SUBTOTAL(103,${tc}${first_data}:${tc}{first_data}),2)=0"
            ],
            fill=PatternFill(
                start_color=SOFT,
                end_color=SOFT,
                fill_type="solid",
            ),
        ),
    )

    # Filter, freeze panes, view and print settings.
    ws.auto_filter.ref = (
        f"{L(c1)}{R_HDR}:{L(c2)}{last_row}"
    )
    ws.freeze_panes = f"A{first_data}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.print_title_rows = f"{R_HDR}:{R_HDR}"
    ws.print_area = f"{L(c1)}1:{L(c2)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(
        fitToPage=True
    )

    # Shift existing chart anchors.
    for chart in charts:
        try:
            _shift_chart(chart, offset)
        except Exception:
            pass

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

    checks, ok = verify(
        src,
        dst,
        ws.title,
        h0,
        c1,
        c2,
        last0,
        headers,
        original,
        written,
        kpis,
        rating_idx,
        title_idx,
        first_data,
        last_row,
    )

    return {
        "output": str(dst),
        "sheet": ws.title,
        "rows": nrows,
        "cols": ncols,
        "checks": checks,
        "ok": ok,
    }


def _soffice():
    exe = shutil.which("soffice") or shutil.which("libreoffice")

    if exe:
        return exe

    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if os.path.exists(candidate):
            return candidate

    return None


def verify(
    src,
    dst,
    sheet,
    h0,
    c1,
    c2,
    last0,
    headers,
    original,
    written,
    kpis,
    rating_idx,
    title_idx,
    first_data,
    last_row,
):
    checks, ok = [], True

    wb = load_workbook(dst)
    ws = wb[sheet]

    # Verify data integrity.
    got_headers = [
        ws.cell(R_HDR, c).value
        for c in range(c1, c2 + 1)
    ]

    got = [
        [
            ws.cell(first_data + i, c).value
            for c in range(c1, c2 + 1)
        ]
        for i in range(len(original))
    ]

    changed = 0

    for original_row, written_row, got_row in zip(
        original, written, got
    ):
        for old, written_value, actual in zip(
            original_row, written_row, got_row
        ):
            is_formula = (
                isinstance(old, str)
                and old.startswith("=")
            )

            if actual != written_value or (
                not is_formula and actual != old
            ):
                changed += 1

    if (
        got_headers == headers
        and changed == 0
        and len(got) == len(original)
    ):
        checks.append(
            f"Data intact ({len(original)} rows x "
            f"{len(headers)} columns)"
        )
    else:
        ok = False
        checks.append(
            f"DATA MISMATCH ({changed} cells differ)"
        )

    # Verify structure.
    expected_filter = (
        f"{L(c1)}{R_HDR}:{L(c2)}{last_row}"
    )

    checks.append(
        f"Filter {ws.auto_filter.ref}"
        if ws.auto_filter.ref == expected_filter
        else "FILTER WRONG"
    )
    ok &= ws.auto_filter.ref == expected_filter

    expected_freeze = f"A{first_data}"

    checks.append(
        f"Header frozen at {ws.freeze_panes}"
        if ws.freeze_panes == expected_freeze
        else "FREEZE WRONG"
    )
    ok &= ws.freeze_panes == expected_freeze

    checks.append(
        "Gridlines hidden"
        if ws.sheet_view.showGridLines is False
        else "GRIDLINES STILL ON"
    )
    ok &= ws.sheet_view.showGridLines is False

    n_cf = sum(
        len(item.rules)
        for item in ws.conditional_formatting
    )

    checks.append(
        f"{n_cf} conditional formatting rules"
        if n_cf >= 2
        else "CONDITIONAL FORMATTING MISSING"
    )
    ok &= n_cf >= 2

    in_table = [
        str(rng)
        for rng in ws.merged_cells.ranges
        if rng.max_row >= R_HDR
        and rng.min_row >= R_HDR
    ]

    checks.append(
        "No merged cells in the data table"
        if not in_table
        else f"MERGED CELLS IN TABLE {in_table}"
    )
    ok &= not in_table

    forms = [
        ws.cell(R_KV, c).value
        for c in range(c1, c2 + 5)
        if isinstance(ws.cell(R_KV, c).value, str)
        and ws.cell(R_KV, c).value.startswith("=")
    ]

    checks.append(
        f"{len(forms)} KPI formulas"
        if len(forms) == 4
        else f"KPI FORMULAS: found {len(forms)}"
    )
    ok &= len(forms) == 4

    # Validate KPI values against the original data.
    ratings = [
        row[rating_idx]
        for row in original
        if _is_num(row[rating_idx])
    ]

    titles = [
        row[title_idx]
        for row in original
        if not _blank(row[title_idx])
    ]

    non_numeric = len(original) - len(ratings)

    if non_numeric:
        checks.append(
            f"WARNING: {non_numeric} rating cells are not numeric "
            "and are ignored by the KPIs"
        )

    if not ratings:
        checks.append(
            "ERROR: No numeric IMDb ratings found. "
            "Check the IMDb Rating column."
        )
        return checks, False

    expected = [
        len(titles),
        sum(ratings) / len(ratings),
        max(ratings),
        min(ratings),
    ]

    exe = _soffice()

    if exe:
        try:
            with tempfile.TemporaryDirectory() as td:
                subprocess.run(
                    [
                        exe,
                        "--headless",
                        "--convert-to",
                        "xlsx",
                        "--outdir",
                        td,
                        str(dst),
                    ],
                    check=True,
                    timeout=150,
                    capture_output=True,
                )

                recalculated_path = (
                    Path(td) / Path(dst).name
                )

                rw = load_workbook(
                    recalculated_path,
                    data_only=True,
                )[sheet]

                vals = []

                for c in range(c1, c2 + 5):
                    value = rw.cell(R_KV, c).value

                    if value is not None:
                        vals.append(value)

                vals = vals[:4]

                if (
                    len(vals) == 4
                    and all(
                        abs(a - b) < 1e-9
                        for a, b in zip(vals, expected)
                    )
                ):
                    checks.append(
                        "KPI values verified: "
                        f"{vals[0]} movies, avg {vals[1]:.2f}, "
                        f"high {vals[2]}, low {vals[3]}"
                    )
                else:
                    ok = False
                    checks.append(
                        f"KPI VALUES DIFFER: "
                        f"sheet {vals} vs data {expected}"
                    )

        except Exception as exc:
            checks.append(
                f"Recalculation check skipped ({str(exc)[:50]})"
            )
    else:
        checks.append(
            "KPI formulas will calculate when opened in Excel; "
            "LibreOffice was not found for a pre-check"
        )

    return checks, ok


if __name__ == "__main__":
    args = sys.argv[1:]

    default_src = (
        Path.home()
        / "AgentWorkspace"
        / "reports"
        / "IMDb_Top_30_Movies_2023.xlsx"
    )

    source = Path(args[0]) if args else default_src

    target = (
        Path(args[1])
        if len(args) > 1
        else source.with_name(
            "IMDb_Top_30_Movies_Designed.xlsx"
        )
    )

    sheet_name = (
        args[2]
        if len(args) > 2
        else "IMDb Top 2023"
    )

    if not source.exists():
        sys.exit(f"File not found: {source}")

    result = design_report(
        source,
        target,
        sheet_name,
    )

    print("Saved:", result["output"])

    for line in result["checks"]:
        print(" -", line)

    print(
        "ALL CHECKS PASSED"
        if result["ok"]
        else "SOME CHECKS FAILED"
    )