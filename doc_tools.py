"""Cross-format document tools: read templates/samples and export to PDF.

Functions: read_document (.docx .xlsx .pptx .pdf .txt .md), export_pdf.
REGISTRY lists (group, schema, function) for main.py.
"""
import os
import shutil
import subprocess

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from common import ROOT, confirm, safe, tool

# ---------- Reading templates / samples ----------
def read_document(path: str) -> str:
    p = safe(path)
    ext = p.suffix.lower()
    if ext == ".docx":
        d = Document(str(p))
        out = []
        for par in d.paragraphs:
            t = par.text.strip()
            if t:
                out.append(f"[{par.style.name}] {t[:200]}")
        out.append(f"(tables: {len(d.tables)})")
        for t in d.tables[:3]:
            out.append("TABLE header: " + " | ".join(c.text.strip() for c in t.rows[0].cells))
        return "\n".join(out)[:4000]
    if ext == ".xlsx":
        wb = load_workbook(str(p), read_only=True)
        out = []
        for ws in wb.worksheets[:4]:
            out.append(f"[Sheet] {ws.title} ({ws.max_row} rows x {ws.max_column} cols)")
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= 6:
                    break
                out.append(" | ".join("" if v is None else str(v) for v in row)[:200])
        return "\n".join(out)[:4000]
    if ext == ".pptx":
        prs = Presentation(str(p))
        out = [f"{len(prs.slides)} slides. Layouts: " +
               ", ".join(l.name for l in prs.slide_layouts)]
        for i, s in enumerate(prs.slides, 1):
            title = s.shapes.title.text if s.shapes.title is not None else ""
            out.append(f"Slide {i}: {title}")
            for sh in s.shapes:
                if sh.has_text_frame and sh != s.shapes.title and sh.text_frame.text.strip():
                    out.append("   - " + sh.text_frame.text.strip().replace("\n", " / ")[:100])
        return "\n".join(out)[:4000]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return "pypdf is not installed. Run: pip install pypdf"
        r = PdfReader(str(p))
        return "\n".join((pg.extract_text() or "") for pg in r.pages[:10])[:4000]
    return p.read_text(encoding="utf-8", errors="replace")[:4000]


# ---------- PDF export ----------
def export_pdf(path: str) -> str:
    p = safe(path)
    ext = p.suffix.lower()
    if ext not in (".docx", ".pptx", ".xlsx") or not p.exists():
        return "Error: give an existing .docx, .pptx or .xlsx file."
    if not confirm(f"Export to PDF: {path}"):
        return "User denied."
    out = p.with_suffix(".pdf")
    rel = lambda f: f.relative_to(ROOT.resolve()).as_posix()
    if ext == ".docx":
        try:
            from docx2pdf import convert
            convert(str(p), str(out))
            if out.exists():
                return f"PDF created: {rel(out)}"
        except Exception:
            pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice") or next(
        (c for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")
         if os.path.exists(c)), None)
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                            "--outdir", str(p.parent), str(p)],
                           check=True, timeout=180, capture_output=True)
            if out.exists():
                return f"PDF created: {rel(out)}"
        except Exception as e:
            return f"Error: PDF export failed: {e}"
    return ("Error: PDF export needs Microsoft Word with the docx2pdf package "
            "(pip install docx2pdf) or LibreOffice installed.")


# ---------- Registry ----------
REGISTRY = [
    ("base", tool("read_document", "Read a sample or template (.docx .xlsx .pptx .pdf .txt .md)",
                  {"path": "e.g. templates/sample.docx"}, ["path"]), read_document),
    ("pdf", tool("export_pdf", "Convert a .docx/.pptx/.xlsx to PDF (needs Word or LibreOffice)",
                 {"path": "The file to convert"}, ["path"]), export_pdf),
]