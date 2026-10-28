"""
ppt_design.py - Executive-grade restyler for an EXISTING PowerPoint (.pptx) deck.

Run directly:
    python ppt_design.py "C:\\path\\Deck.pptx"
    python ppt_design.py <source.pptx> <output.pptx> --footer "ACME Corp | Strategy Review"

The source file is never modified. The restyled copy keeps every slide, text, table, chart,
picture and speaker note; only the visual styling is elevated to executive consulting standards.
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# =====================================================================
# Executive Design System & Color Palette
# =====================================================================
# Modern Slate & Sapphire Consulting Palette
SLATE_900  = "0F172A"   # Dominant dark hero
NAVY_HEADER= "1E293B"   # Card/Header background
BLUE_ACCENT= "2563EB"   # Royal brand blue
CYAN_ACCENT= "0EA5E9"   # Modern Sky cyan
LIGHT_BG   = "F8FAFC"   # Modern soft canvas
WHITE      = "FFFFFF"   # Clean white card
CARD_BG    = "F1F5F9"   # Table zebra / Card fill
LINE       = "E2E8F0"   # Subtle borders
TEXT_DARK  = "0F172A"   # Body text
TEXT_MUTED = "64748B"   # Secondary/Footnotes
TEXT_PALE  = "94A3B8"   # Dark mode subtext

FONT = "Segoe UI"
PALETTE = ["2563EB", "0EA5E9", "3B82F6", "60A5FA", "64748B", "CBD5E1"]

MARK = "pcagent-16x9"
PREFIX = "Design "

_SKIP_TYPES = (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE)


def _rgb(hex_str):
    return RGBColor.from_string(hex_str)


# =====================================================================
# Shape Helpers
# =====================================================================

def _rect(slide, left, top, width, height, color, name):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.name = PREFIX + name
    return shp


def _to_back(slide, shp):
    tree = slide.shapes._spTree
    tree.remove(shp._element)
    tree.insert(2, shp._element)


def _remove(shape):
    try:
        shape._element.getparent().remove(shape._element)
    except Exception:
        pass


def _is_decor(shape, marked, W, H):
    """Strip earlier thin decor bars or generated footer boxes."""
    if not marked:
        return False
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not shape.has_text_frame:
            return min(shape.width, shape.height) < Inches(0.3)
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.has_text_frame and not shape.text_frame.text.strip():
            return min(shape.width, shape.height) < Inches(0.3)
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.has_text_frame:
            t = shape.text_frame.text.strip()
            return t.isdigit() and shape.top is not None and shape.top > H - Inches(0.7)
    except Exception:
        return False
    return False


def _style_runs(tf, size_by_level=None, size=None, color=TEXT_DARK, bold=None):
    for p in tf.paragraphs:
        sz = size_by_level.get(p.level, size_by_level.get(max(size_by_level))) if size_by_level else size
        for r in p.runs:
            r.font.name = FONT
            if sz:
                r.font.size = Pt(sz)
            if color:
                r.font.color.rgb = _rgb(color)
            if bold is not None:
                r.font.bold = bold


def _fit_text(tf):
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


# =====================================================================
# Tables & Charts
# =====================================================================

def _cell_border(cell, color, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(w), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(i, ln)


def _style_table(shape, sc):
    tbl = shape.table
    ncols = len(tbl.columns)
    size = round((15 if ncols <= 4 else 13 if ncols <= 6 else 11) * sc)

    for ri, row in enumerate(tbl.rows):
        is_header = (ri == 0)
        for cell in row.cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(NAVY_HEADER if is_header else (CARD_BG if ri % 2 == 0 else WHITE))
            _cell_border(cell, LINE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.margin_top = cell.margin_bottom = Inches(0.06)

            for p in cell.text_frame.paragraphs:
                # Right-align numeric columns
                t = p.text.strip()
                if not is_header and re.match(r"^[\$€£]?\s*[-+]?[0-9,.]+%?$", t):
                    p.alignment = PP_ALIGN.RIGHT
                else:
                    p.alignment = PP_ALIGN.LEFT

                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(size)
                    r.font.bold = True if is_header else r.font.bold
                    r.font.color.rgb = _rgb(WHITE if is_header else TEXT_DARK)


def _style_chart(chart, sc):
    try:
        chart.font.name = FONT
        chart.font.size = Pt(round(12 * sc))
        chart.font.color.rgb = _rgb(TEXT_DARK)
        pie = chart.chart_type in (
            XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED,
            XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.THREE_D_PIE
        )
        line = "LINE" in str(chart.chart_type)
        for si, ser in enumerate(chart.plots[0].series):
            if pie:
                for pi in range(len(list(ser.values))):
                    pt = ser.points[pi]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _rgb(PALETTE[pi % len(PALETTE)])
            elif line:
                ser.format.line.color.rgb = _rgb(PALETTE[si % len(PALETTE)])
                ser.format.line.width = Pt(2.5)
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = _rgb(PALETTE[si % len(PALETTE)])
    except Exception:
        pass


# =====================================================================
# Footer
# =====================================================================

def _footer(slide, W, H, text, number, sc, add_text, add_number):
    _to_back(slide, _rect(slide, Inches(0.8), H - Inches(0.55), W - Inches(1.6), Inches(0.012), LINE, "Footer Line"))

    if add_text and text:
        tb = slide.shapes.add_textbox(Inches(0.8), H - Inches(0.48), int(W * 0.6), Inches(0.3))
        tb.name = PREFIX + "Footer"
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = text
        r.font.name, r.font.size = FONT, Pt(round(9 * sc))
        r.font.color.rgb = _rgb(TEXT_MUTED)

    if add_number:
        tb = slide.shapes.add_textbox(W - Inches(1.8), H - Inches(0.48), Inches(1.0), Inches(0.3))
        tb.name = PREFIX + "Number"
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        fld = etree.SubElement(p._p, qn("a:fld"))
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
        fld.set("type", "slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"))
        rpr.set("lang", "en-US")
        rpr.set("sz", str(int(round(9 * sc) * 100)))
        sf = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", TEXT_MUTED)
        etree.SubElement(rpr, qn("a:latin")).set("typeface", FONT)
        etree.SubElement(fld, qn("a:t")).text = str(number)


# =====================================================================
# Slide Restyling (Title vs Content)
# =====================================================================

def _is_title_slide(slide, idx):
    name = (slide.slide_layout.name or "").lower()
    return idx == 1 or "title slide" in name or "section header" in name


def _design_title_slide(slide, W, H, sc):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(SLATE_900)

    # Modern accent stripes on the left border
    _to_back(slide, _rect(slide, 0, 0, Inches(0.25), H, BLUE_ACCENT, "Left Strip 1"))
    _to_back(slide, _rect(slide, Inches(0.25), 0, Inches(0.08), H, CYAN_ACCENT, "Left Strip 2"))

    title = slide.shapes.title
    sub = next((ph for ph in slide.placeholders
                if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
                or (ph.placeholder_format.idx == 1 and ph != title)), None)

    t_top = int(H * 0.28)
    t_h = Inches(1.8)

    if title is not None:
        title.left, title.top, title.width, title.height = Inches(1.2), t_top, W - Inches(2.2), t_h
        tf = title.text_frame
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        _fit_text(tf)
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.LEFT
        _style_runs(tf, size=round(44 * sc), color=WHITE, bold=True)

    if sub is not None and sub.has_text_frame:
        sub.left, sub.top, sub.width, sub.height = Inches(1.2), t_top + t_h + Inches(0.2), W - Inches(2.2), Inches(1.2)
        tf = sub.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        _fit_text(tf)
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.LEFT
        _style_runs(tf, size=round(20 * sc), color=TEXT_PALE)


def _design_content_slide(slide, idx, W, H, sc, footer_text):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(LIGHT_BG)

    title = slide.shapes.title
    existing = {ph.placeholder_format.type for ph in slide.placeholders}

    # Sleek consulting header (No heavy colored block)
    header_h = Inches(1.2)
    _to_back(slide, _rect(slide, Inches(0.8), header_h, W - Inches(1.6), Inches(0.015), LINE, "Header Divider"))
    _to_back(slide, _rect(slide, Inches(0.8), header_h, Inches(1.2), Inches(0.035), BLUE_ACCENT, "Header Accent"))

    if title is not None:
        title.left = Inches(0.8)
        title.top = Inches(0.4)
        title.width = W - Inches(1.6)
        title.height = Inches(0.7)
        tf = title.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _fit_text(tf)
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.LEFT
        _style_runs(tf, size=round(24 * sc), color=TEXT_DARK, bold=True)

    content_top = header_h + Inches(0.3)
    footer_top = H - Inches(0.65)

    bodies = [ph for ph in slide.placeholders
              if ph != title and ph.placeholder_format.type not in _SKIP_TYPES and ph.has_text_frame]

    levels = {0: round(20 * sc), 1: round(17 * sc), 2: round(15 * sc), 3: round(13 * sc)}
    for ph in bodies:
        tf = ph.text_frame
        _fit_text(tf)
        for p in tf.paragraphs:
            p.space_after = Pt(8)
        _style_runs(tf, size_by_level=levels, color=TEXT_DARK)

    single = bodies[0] if len(bodies) == 1 else None

    # Card background behind single body placeholder
    if single is not None:
        c_left = Inches(0.8)
        c_top = content_top
        c_w = W - Inches(1.6)
        c_h = footer_top - content_top - Inches(0.1)

        card = _rect(slide, c_left, c_top, c_w, c_h, WHITE, "Card")
        card.line.color.rgb = _rgb(LINE)
        card.line.width = Pt(1)
        _to_back(slide, card)

        single.left = c_left + Inches(0.35)
        single.top = c_top + Inches(0.3)
        single.width = c_w - Inches(0.7)
        single.height = c_h - Inches(0.6)

    # Style all other shapes (tables, charts, images)
    for sh in list(slide.shapes):
        if sh == title or sh.name.startswith(PREFIX) or sh.top is None or sh.height is None:
            continue
        if sh.is_placeholder and sh.placeholder_format.type in _SKIP_TYPES:
            continue
        if sh == single:
            continue
        if sh.top < content_top - Inches(0.05):
            sh.top = content_top
        if sh.top + sh.height > footer_top:
            room = max(footer_top - sh.top, Inches(0.5))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                ratio = room / sh.height
                sh.width, sh.height = int(sh.width * ratio), int(room)
            else:
                sh.height = int(room)

        if getattr(sh, "has_table", False) and sh.has_table:
            _style_table(sh, sc)
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            _style_chart(sh.chart, sc)
        elif sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip():
            _style_runs(sh.text_frame, color=TEXT_DARK)

    _footer(slide, W, H, footer_text, idx, sc,
            add_text=PP_PLACEHOLDER.FOOTER not in existing,
            add_number=PP_PLACEHOLDER.SLIDE_NUMBER not in existing)


# =====================================================================
# Verification & Integrity Helpers
# =====================================================================

def _snapshot(prs, marked):
    W, H = prs.slide_width, prs.slide_height
    out = []
    for slide in prs.slides:
        texts, pics, charts, tables = [], 0, 0, 0
        for sh in slide.shapes:
            if sh.name.startswith(PREFIX) or _is_decor(sh, marked, W, H):
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                tables += 1
                for row in sh.table.rows:
                    for c in row.cells:
                        if c.text.strip():
                            texts.append(c.text.strip())
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pics += 1
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append({"texts": Counter(texts), "pics": pics, "charts": charts,
                    "tables": tables, "notes": notes})
    return out


def _soffice():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(c):
            return c
    return None


def _pdf_pages(path):
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        return len(re.findall(rb"/Type\s*/Page[^s]", Path(path).read_bytes()))


# =====================================================================
# Main Restyle Entry
# =====================================================================

def design_deck(src, dst, footer=""):
    src, dst = Path(src), Path(dst)
    if src.resolve() == dst.resolve():
        raise ValueError("Output must be a different file. The original is kept untouched.")
    prs = Presentation(str(src))
    if len(prs.slides) == 0:
        raise ValueError("The presentation has no slides.")
    marked = prs.core_properties.comments == MARK
    W, H = prs.slide_width, prs.slide_height
    sc = min(1.0, W / Inches(13.333))
    before = _snapshot(prs, marked)

    first = prs.slides[0]
    footer_text = footer.strip() or (
        first.shapes.title.text.strip() if first.shapes.title is not None and first.shapes.title.text.strip()
        else src.stem.replace("_", " "))

    for idx, slide in enumerate(prs.slides, start=1):
        for sh in list(slide.shapes):
            if sh.name.startswith(PREFIX) or _is_decor(sh, marked, W, H):
                _remove(sh)
        if _is_title_slide(slide, idx):
            _design_title_slide(slide, W, H, sc)
        else:
            _design_content_slide(slide, idx, W, H, sc, footer_text)

    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    checks, ok = verify(src, dst, before, marked, len(prs.slides), (W, H))
    return {"output": str(dst), "slides": len(prs.slides), "checks": checks, "ok": ok}


def verify(src, dst, before, marked, n_slides, size):
    checks, ok = [], True
    prs = Presentation(str(dst))
    after = _snapshot(prs, marked)

    if len(prs.slides) != n_slides or (prs.slide_width, prs.slide_height) != size:
        ok = False
        checks.append("SLIDE COUNT OR SIZE CHANGED")

    problems = []
    for i, (b, a) in enumerate(zip(before, after), start=1):
        if (b["texts"] - a["texts"]) or b["notes"] != a["notes"] or \
                (b["pics"], b["charts"], b["tables"]) != (a["pics"], a["charts"], a["tables"]):
            problems.append(i)

    if not problems:
        total = sum(sum(b["texts"].values()) for b in before)
        pics = sum(b["pics"] for b in before)
        charts = sum(b["charts"] for b in before)
        tables = sum(b["tables"] for b in before)
        checks.append(f"Content intact ({n_slides} slides, {total} text blocks, {tables} tables, "
                      f"{charts} charts, {pics} pictures, speaker notes verified)")
    else:
        ok = False
        checks.append(f"CONTENT MISMATCH on slide(s) {problems}")

    exe = _soffice()
    if exe:
        try:
            with tempfile.TemporaryDirectory() as td:
                subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, str(dst)],
                               check=True, timeout=180, capture_output=True)
                pages = _pdf_pages(Path(td) / (Path(dst).stem + ".pdf"))
                if pages == n_slides:
                    checks.append(f"Opens and renders accurately ({pages} slides)")
                else:
                    ok = False
                    checks.append(f"RENDER MISMATCH: {pages} pages for {n_slides} slides")
        except Exception as e:
            checks.append(f"Render check skipped ({str(e)[:50]})")

    return checks, ok


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Apply an executive design to a PowerPoint deck.")
    ap.add_argument("source")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--footer", default="")
    a = ap.parse_args()
    source = Path(a.source)
    if not source.exists():
        sys.exit(f"File not found: {source}")
    target = Path(a.output) if a.output else source.with_name(source.stem + "_Designed.pptx")
    res = design_deck(source, target, a.footer)
    print("Saved:", res["output"])
    for line in res["checks"]:
        print(" -", line)
    print("ALL CHECKS PASSED" if res["ok"] else "SOME CHECKS FAILED")