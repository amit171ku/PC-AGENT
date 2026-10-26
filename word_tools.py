"""Word (.docx) tools: build and edit professional reports.

Functions: report_start, report_add_content, report_add_chart, report_add_image, replace_text.
REGISTRY lists (group, schema, function) for main.py.
"""
import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from common import (approved, confirm, guard, parse_floats, parse_pipe_rows, safe,
                    split_labels, tool)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT


# ---------- Word (.docx) ----------
def _font(style, name="Calibri", size=None, bold=None, color=None):
    style.font.name = name
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rf = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]


def _new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)
    _font(doc.styles["Normal"], size=11)
    pf = doc.styles["Normal"].paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for name, size in (("Title", 30), ("Subtitle", 16), ("Heading 1", 16),
                       ("Heading 2", 13), ("Heading 3", 12)):
        _font(doc.styles[name], size=size, bold=(name != "Subtitle"),
              color=(GREY if name == "Subtitle" else NAVY))
    return doc


def _clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _field(par, instr, placeholder=""):
    def mk(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    par.add_run()._r.append(mk("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    par.add_run()._r.append(it)
    par.add_run()._r.append(mk("separate"))
    par.add_run(placeholder)
    par.add_run()._r.append(mk("end"))


def _styled(doc, text, style, size):
    try:
        return doc.add_paragraph(text, style=style)
    except KeyError:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        return p


def _header_footer(doc, title, company, cover):
    for s in doc.sections:
        if cover:
            s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]
        hp.alignment = RIGHT
        r = hp.add_run(" | ".join(x for x in (company, title) if x))
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        fp = s.footer.paragraphs[0]
        fp.alignment = CENTER
        r = fp.add_run("Page ")
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        _field(fp, "PAGE", "1")


def _cover(doc, title, subtitle, company, author, date):
    for _ in range(5):
        doc.add_paragraph()
    _styled(doc, title, "Title", 28).alignment = CENTER
    if subtitle:
        _styled(doc, subtitle, "Subtitle", 16).alignment = CENTER
    doc.add_paragraph()
    for line in (company, author, date):
        if line:
            p = doc.add_paragraph()
            p.alignment = CENTER
            p.add_run(line).font.size = Pt(12)
    doc.add_page_break()


def _toc(doc):
    p = doc.add_paragraph()
    r = p.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = NAVY
    _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
           "Open in Word and press F9 (or accept the update prompt) to build the table of contents.")
    doc.add_page_break()
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)


def report_start(path: str, title: str, subtitle: str = "", company: str = "",
                 author: str = "", date: str = "", template: str = "",
                 include_cover: str = "yes", include_toc: str = "yes") -> str:
    p = safe(path)
    if p.suffix.lower() != ".docx":
        return "Error: path must end with .docx"
    note = f" using template {template}" if template else ""
    if not confirm(f"Create Word report: {path}{note}"):
        return "User denied."
    cover = include_cover.strip().lower() != "no"
    toc = include_toc.strip().lower() != "no"
    if template:
        t = safe(template)
        if t.suffix.lower() != ".docx":
            return "Error: template must be a .docx file."
        doc = Document(str(t))
        _clear_body(doc)
    else:
        doc = _new_doc()
        _header_footer(doc, title, company, cover)
    if cover:
        _cover(doc, title, subtitle, company, author,
               date or datetime.now().strftime("%d %B %Y"))
    if toc:
        _toc(doc)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(p))
    approved.add(str(p))
    return f"Word report created at {path}. Add content with report_add_content."


def _runs(par, text):
    for tok in re.split(r"(\*\*[^*]+\*\*)", text):
        if not tok:
            continue
        if len(tok) > 4 and tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        else:
            par.add_run(tok)


def _heading(doc, text, level):
    try:
        level = max(1, min(int(level or 1), 3))
    except ValueError:
        level = 1
    try:
        return doc.add_heading(text, level=level)
    except KeyError:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt({1: 16, 2: 13, 3: 12}[level])
        return p


def _shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc.append(shd)


def _table(doc, rows):
    if not rows:
        return
    n = len(rows[0])
    table = doc.add_table(rows=1, cols=n)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = WHITE
        _shade(cell, "1F3A5F")
    for row in rows[1:]:
        cells = table.add_row().cells
        for i in range(n):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(row[i] if i < len(row) else "").font.size = Pt(10)
    doc.add_paragraph()


def _add_markdown(doc, md):
    bullet = re.compile(r"^\s*[-*\u2022]\s+")
    number = re.compile(r"^\s*\d+[.)]\s+")
    lines = md.replace("\r", "").split("\n")
    para = []

    def flush():
        if para:
            _runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            flush()
            _heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue
        if s.lower() in ("[pagebreak]", "---pagebreak---"):
            flush()
            doc.add_page_break()
            i += 1
            continue
        if s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, parse_pipe_rows("\n".join(block)))
            continue
        if bullet.match(s):
            flush()
            try:
                par = doc.add_paragraph(style="List Bullet")
            except KeyError:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Inches(0.25)
                par.add_run("\u2022 ")
            _runs(par, bullet.sub("", s))
            i += 1
            continue
        if number.match(s):
            flush()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.25)
            _runs(par, s)
            i += 1
            continue
        para.append(s)
        i += 1
    flush()


def report_add_content(path: str, content: str) -> str:
    p = guard(path, ".docx")
    doc = Document(str(p))
    _add_markdown(doc, content)
    doc.save(str(p))
    return f"Added {len(content.split())} words to {path}."


def report_add_chart(path: str, chart_type: str, title: str, labels: str,
                     values: str, series_name: str = "", caption: str = "") -> str:
    try:
        from matplotlib.figure import Figure
    except ImportError:
        return "Error: matplotlib is not installed. Run: pip install matplotlib"
    p = guard(path, ".docx")
    doc = Document(str(p))
    labs, vals = split_labels(labels), parse_floats(values)
    n = min(len(labs), len(vals))
    labs, vals = labs[:n], vals[:n]
    t = chart_type.strip().lower()
    fig = Figure(figsize=(6.5, 3.4), dpi=150)
    ax = fig.subplots()
    if t == "pie":
        ax.pie(vals, labels=labs, autopct="%1.0f%%", startangle=90)
        ax.axis("equal")
    elif t == "line":
        ax.plot(labs, vals, marker="o", color="#1F3A5F")
        ax.grid(alpha=0.3)
    else:
        ax.bar(labs, vals, color="#1F3A5F")
    if t != "pie":
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        if series_name:
            ax.set_ylabel(series_name)
    ax.set_title(title, fontsize=11, fontweight="bold")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    buf.seek(0)
    doc.add_picture(buf, width=Inches(6))
    doc.paragraphs[-1].alignment = CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = CENTER
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
    doc.save(str(p))
    return "Chart added."


def report_add_image(path: str, image: str, caption: str = "", width_inches: str = "6") -> str:
    p = guard(path, ".docx")
    img = safe(image)
    if img.suffix.lower() not in (".png", ".jpg", ".jpeg") or not img.exists():
        return "Error: image must be an existing .png or .jpg in the workspace."
    try:
        w = float(width_inches)
    except ValueError:
        w = 6.0
    doc = Document(str(p))
    doc.add_picture(str(img), width=Inches(min(max(w, 1), 6.5)))
    doc.paragraphs[-1].alignment = CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = CENTER
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
    doc.save(str(p))
    return "Image added."


def _iter_pars(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                yield p


def replace_text(path: str, find: str, replace: str) -> str:
    p = guard(path, ".docx")
    doc = Document(str(p))
    count = 0
    for par in _iter_pars(doc):
        if find not in par.text:
            continue
        done = False
        for r in par.runs:
            if find in r.text:
                r.text = r.text.replace(find, replace)
                done = True
        if not done and par.runs:
            par.runs[0].text = par.text.replace(find, replace)
            for r in par.runs[1:]:
                r.text = ""
        count += 1
    doc.save(str(p))
    return f"Replaced in {count} paragraph(s)."


# ---------- Registry ----------
REGISTRY = [
    ("word", tool("report_start", "Create a Word report with optional cover and table of contents. Call first.",
                  {"path": "e.g. reports/Project_Report.docx", "title": "", "subtitle": "", "company": "",
                   "author": "", "date": "",
                   "template": "Optional .docx template to reuse its styles",
                   "include_cover": "yes or no", "include_toc": "yes or no"},
                  ["path", "title"]), report_start),
    ("word", tool("report_add_content",
                  "Append content to a Word report. Markdown: # ## ### headings, - bullets, 1. lists, | tables |, **bold**, [pagebreak]. Several sections per call.",
                  {"path": "The .docx path", "content": "Markdown text"}, ["path", "content"]),
     report_add_content),
    ("word", tool("report_add_chart", "Add a chart image to a Word report",
                  {"path": "The .docx path", "chart_type": "bar, line or pie", "title": "",
                   "labels": "A|B|C", "values": "1|2|3", "series_name": "", "caption": ""},
                  ["path", "chart_type", "title", "labels", "values"]), report_add_chart),
    ("word", tool("report_add_image", "Add an uploaded image to a Word report",
                  {"path": "The .docx path", "image": "e.g. templates/logo.png", "caption": "",
                   "width_inches": "default 6"}, ["path", "image"]), report_add_image),
    ("word", tool("replace_text", "Find and replace text in a .docx (template placeholders)",
                  {"path": "The .docx path", "find": "", "replace": ""},
                  ["path", "find", "replace"]), replace_text),
]