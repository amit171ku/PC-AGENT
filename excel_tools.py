"""Excel (.xlsx) tools: create, read, format, conditional formatting, layout, charts,
and the one-call premium redesign (excel_design.py).

REGISTRY lists (group, schema, function) for main.py.
"""
import re
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

from common import ROOT, approved, confirm, guard, parse_pipe_rows, safe, tool
from excel_design import design_report

# ---------- Excel (.xlsx) ----------
XL_HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
_THIN = Side(style="thin", color="D0D5DD")
XL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _num(v: str):
    s = str(v).strip()
    if s.startswith("="):
        return s
    t = s.replace(",", "")
    if re.fullmatch(r"-?([1-9]\d*|0)", t):
        return int(t)
    if re.fullmatch(r"-?\d*\.\d+", t):
        return float(t)
    return s


def excel_start(path: str, template: str = "") -> str:
    p = safe(path)
    if p.suffix.lower() != ".xlsx":
        return "Error: path must end with .xlsx"
    note = f" using template {template}" if template else ""
    if not confirm(f"Create Excel workbook: {path}{note}"):
        return "User denied."
    p.parent.mkdir(parents=True, exist_ok=True)
    if template:
        t = safe(template)
        if t.suffix.lower() != ".xlsx":
            return "Error: template must be an .xlsx file."
        shutil.copy(t, p)
    else:
        Workbook().save(str(p))
    approved.add(str(p))
    return f"Excel workbook created at {path}. Add tables with excel_add_sheet."


def excel_add_sheet(path: str, sheet: str, table: str, title: str = "") -> str:
    p = guard(path, ".xlsx")
    rows = parse_pipe_rows(table)
    if len(rows) < 2:
        raise ValueError("table needs a header line and at least one data row")
    wb = load_workbook(str(p))
    name = re.sub(r"[\\/*?:\[\]]", "", sheet)[:31] or "Sheet1"
    if name in wb.sheetnames:
        ws = wb[name]
    elif wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        ws = wb["Sheet"]
        ws.title = name
    else:
        ws = wb.create_sheet(name)
    empty = ws.max_row == 1 and ws["A1"].value is None
    r0 = 1 if empty else ws.max_row + 2
    if title:
        ws.cell(r0, 1, title).font = Font(bold=True, size=14, color="1F3A5F")
        r0 += 1
    hdr = r0
    ncols = len(rows[0])
    plain = [bool(re.search(r"year|rank|\bid\b|code|zip|phone|\bno\b|#", h, re.I)) for h in rows[0]]
    for j, h in enumerate(rows[0], 1):
        c = ws.cell(hdr, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = XL_HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = XL_BORDER
    for i, row in enumerate(rows[1:], hdr + 1):
        for j in range(ncols):
            v = _num(row[j]) if j < len(row) else ""
            c = ws.cell(i, j + 1, v)
            c.border = XL_BORDER
            if isinstance(v, int) and not plain[j]:
                c.number_format = "#,##0"
            elif isinstance(v, float):
                c.number_format = "#,##0.00"
    last = hdr + len(rows) - 1
    for j in range(1, ncols + 1):
        letter = get_column_letter(j)
        longest = max(len(str(ws.cell(r, j).value or "")) for r in range(hdr, last + 1))
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 min(max(longest + 2, 10), 50))
    endcol = get_column_letter(ncols)
    if empty:
        ws.freeze_panes = ws.cell(hdr + 1, 1)
        ws.auto_filter.ref = f"A{hdr}:{endcol}{last}"
    wb.save(str(p))
    return (f"Sheet '{ws.title}': header row {hdr}, data rows {hdr + 1}-{last}, "
            f"columns A-{endcol}. Table range A{hdr}:{endcol}{last}.")


def excel_set_cells(path: str, sheet: str, cells: str) -> str:
    p = guard(path, ".xlsx")
    wb = load_workbook(str(p))
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet}. Sheets: {', '.join(wb.sheetnames)}")
    ws = wb[sheet]
    count = 0
    for line in cells.replace("\r", "").split("\n"):
        if "=" not in line:
            continue
        ref, val = line.split("=", 1)
        ref = ref.strip()
        if not re.fullmatch(r"[A-Za-z]{1,3}\d{1,6}", ref):
            continue
        v = _num(val)
        c = ws[ref.upper()]
        c.value = v
        if isinstance(v, str) and v.startswith("="):
            c.font = Font(bold=True)
            c.number_format = "#,##0.00"
        count += 1
    wb.save(str(p))
    return f"Set {count} cell(s). Excel calculates formulas when the file is opened."


def excel_add_chart(path: str, sheet: str, chart_type: str, title: str,
                    data_range: str, categories_range: str, anchor: str = "H2") -> str:
    p = guard(path, ".xlsx")
    wb = load_workbook(str(p))
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet}. Sheets: {', '.join(wb.sheetnames)}")
    ws = wb[sheet]
    kinds = {"bar": BarChart, "line": LineChart, "pie": PieChart}
    chart = kinds.get(chart_type.strip().lower(), BarChart)()
    if isinstance(chart, BarChart):
        chart.type = "col"
    chart.title = title
    data = Reference(ws, range_string=f"'{ws.title}'!{data_range}")
    cats = Reference(ws, range_string=f"'{ws.title}'!{categories_range}")
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width, chart.height = 18, 9
    ws.add_chart(chart, anchor)
    wb.save(str(p))
    return "Chart added."



# ---- Excel: read, format, conditional formatting, layout ----
_COLORS = {
    "red": "FF0000", "green": "00B050", "blue": "0070C0", "yellow": "FFFF00",
    "orange": "FFA500", "grey": "808080", "gray": "808080", "white": "FFFFFF",
    "black": "000000", "navy": "1F3A5F", "purple": "7030A0", "gold": "FFC000",
    "lightgreen": "C6EFCE", "lightred": "FFC7CE", "lightyellow": "FFEB9C",
    "lightblue": "DDEBF7", "lightgrey": "F2F2F2", "lightgray": "F2F2F2",
}


def _hex(c, default="000000"):
    c = (c or "").strip().lstrip("#").lower()
    if c in _COLORS:
        return _COLORS[c]
    if re.fullmatch(r"[0-9a-f]{6}", c):
        return c.upper()
    return default


def _kv(text):
    out = {}
    for part in re.split(r"[|\n]", text or ""):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _bool(v):
    return str(v).strip().lower() in ("1", "true", "yes", "on")


def _flat(x):
    if isinstance(x, tuple):
        for i in x:
            yield from _flat(i)
    else:
        yield x


def _xl_sheet(wb, sheet):
    if sheet in wb.sheetnames:
        return wb[sheet]
    if not sheet and wb.worksheets:
        return wb.worksheets[0]
    raise ValueError(f"Sheet not found: {sheet}. Sheets: {', '.join(wb.sheetnames)}")


def excel_read(path: str, sheet: str = "", cell_range: str = "") -> str:
    p = safe(path)
    if p.suffix.lower() != ".xlsx" or not p.exists():
        raise ValueError(f"File not found: {path}")
    wb = load_workbook(str(p))
    ws = _xl_sheet(wb, sheet)
    rng = cell_range.strip() or (
        f"A1:{get_column_letter(min(ws.max_column, 12))}{min(ws.max_row, 25)}")
    head = (f"Sheets: {', '.join(wb.sheetnames)}. '{ws.title}' is {ws.max_row} rows x "
            f"{ws.max_column} cols. Range {rng}:")
    lines = [head]
    for row in ws[rng.replace("$", "")]:
        cells = list(_flat(row))
        if not cells:
            continue
        vals = ["" if c.value is None else str(c.value)[:40] for c in cells]
        lines.append(f"r{cells[0].row}: " + " | ".join(vals))
        if sum(len(x) for x in lines) > 3500:
            lines.append("...(truncated)")
            break
    return "\n".join(lines)


_PRESETS = {
    "header": "bold=true|color=FFFFFF|fill=1F3A5F|align=center|valign=center|wrap=true|border=thin",
    "title": "bold=true|size=14|color=1F3A5F",
    "total": "bold=true|fill=E8EEF7|border=thin",
    "currency": "number_format=#,##0.00",
    "integer": "number_format=#,##0",
    "percent": "number_format=0.0%",
    "date": "number_format=dd-mmm-yyyy",
}


def _border(spec, color):
    name, _, sides = spec.partition(":")
    name = name.strip().lower()
    if name in ("none", ""):
        return Border()
    side = Side(style=name, color=_hex(color, "BFBFBF"))
    want = {x.strip() for x in (sides or "left,right,top,bottom").split(",")}
    return Border(left=side if "left" in want else Side(),
                  right=side if "right" in want else Side(),
                  top=side if "top" in want else Side(),
                  bottom=side if "bottom" in want else Side())


def excel_format(path: str, sheet: str, cell_range: str, style: str) -> str:
    p = guard(path, ".xlsx")
    wb = load_workbook(str(p))
    ws = _xl_sheet(wb, sheet)
    st = {}
    for part in re.split(r"[|\n]", style):
        part = part.strip().lower()
        if part.startswith("preset="):
            part = part[7:].strip()
        if part in _PRESETS:
            st.update(_kv(_PRESETS[part]))
    st.update({k: v for k, v in _kv(style).items() if k != "preset"})
    if "preset" in _kv(style):
        base = _kv(_PRESETS.get(_kv(style)["preset"].lower(), ""))
        st = {**base, **{k: v for k, v in _kv(style).items() if k != "preset"}}
    count = 0
    for c in _flat(ws[cell_range.strip().replace("$", "")]):
        if count >= 20000:
            break
        f = c.font
        c.font = Font(
            name=st.get("font", f.name),
            size=float(st["size"]) if "size" in st else f.size,
            bold=_bool(st["bold"]) if "bold" in st else f.bold,
            italic=_bool(st["italic"]) if "italic" in st else f.italic,
            underline="single" if _bool(st.get("underline", "")) else f.underline,
            color=_hex(st["color"]) if "color" in st else f.color,
        )
        if "fill" in st:
            c.fill = (PatternFill(fill_type=None) if st["fill"].lower() == "none"
                      else PatternFill("solid", fgColor=_hex(st["fill"])))
        a = c.alignment
        if any(k in st for k in ("align", "valign", "wrap", "indent")):
            c.alignment = Alignment(
                horizontal=st.get("align", a.horizontal),
                vertical=st.get("valign", a.vertical),
                wrap_text=_bool(st["wrap"]) if "wrap" in st else a.wrap_text,
                indent=int(st["indent"]) if "indent" in st else a.indent)
        if "border" in st:
            c.border = _border(st["border"], st.get("border_color", "BFBFBF"))
        if "number_format" in st:
            c.number_format = st["number_format"]
        count += 1
    wb.save(str(p))
    return f"Formatted {count} cell(s) in {cell_range}."


def excel_conditional_format(path: str, sheet: str, cell_range: str, rule: str,
                             value: str = "", fill: str = "", font_color: str = "",
                             colors: str = "") -> str:
    p = guard(path, ".xlsx")
    wb = load_workbook(str(p))
    ws = _xl_sheet(wb, sheet)
    rng = cell_range.strip().replace("$", "")
    first = rng.split(":")[0]
    r = rule.strip().lower().replace(" ", "_")
    zebra = r == "zebra"
    if not fill and not font_color:
        fill, font_color = ("F2F2F2", "") if zebra else ("FFC7CE", "9C0006")
    fobj = PatternFill(start_color=_hex(fill), end_color=_hex(fill), fill_type="solid") if fill else None
    font = Font(color=_hex(font_color)) if font_color else None
    kw = {k: v for k, v in (("fill", fobj), ("font", font)) if v is not None}

    def lit(v):
        v = v.strip()
        return v if re.fullmatch(r"-?\d+(\.\d+)?", v) else '"' + v.replace('"', '""') + '"'

    ops = {"greater_than": "greaterThan", "less_than": "lessThan", "equal": "equal",
           "not_equal": "notEqual", "at_least": "greaterThanOrEqual", "at_most": "lessThanOrEqual"}
    if r == "color_scale":
        cols = [_hex(x) for x in colors.split("|") if x.strip()] or ["F8696B", "FFEB84", "63BE7B"]
        if len(cols) == 2:
            obj = ColorScaleRule(start_type="min", start_color=cols[0],
                                 end_type="max", end_color=cols[1])
        else:
            obj = ColorScaleRule(start_type="min", start_color=cols[0],
                                 mid_type="percentile", mid_value=50, mid_color=cols[1],
                                 end_type="max", end_color=cols[-1])
    elif r == "data_bar":
        obj = DataBarRule(start_type="min", end_type="max",
                          color=_hex(colors or fill, "638EC6"), showValue=True)
    elif r in ops:
        obj = CellIsRule(operator=ops[r], formula=[lit(value)], **kw)
    elif r in ("between", "not_between"):
        a, _, b = value.partition("|")
        obj = CellIsRule(operator="between" if r == "between" else "notBetween",
                         formula=[lit(a), lit(b)], **kw)
    elif r == "text_contains":
        obj = FormulaRule(formula=[f'ISNUMBER(SEARCH({lit(value)},{first}))'], **kw)
    elif r in ("top_n", "bottom_n"):
        obj = Rule(type="top10", rank=int(float(value or 10)), bottom=(r == "bottom_n"),
                   dxf=DifferentialStyle(**kw))
    elif r == "duplicates":
        obj = Rule(type="duplicateValues", dxf=DifferentialStyle(**kw))
    elif r == "zebra":
        obj = FormulaRule(formula=["MOD(ROW(),2)=0"], **kw)
    elif r == "formula":
        obj = FormulaRule(formula=[value.lstrip("=")], **kw)
    else:
        raise ValueError("Unknown rule. Use: color_scale, data_bar, greater_than, less_than, equal, "
                         "not_equal, at_least, at_most, between, text_contains, top_n, bottom_n, "
                         "duplicates, zebra, formula")
    ws.conditional_formatting.add(rng, obj)
    wb.save(str(p))
    return f"Conditional formatting '{r}' applied to {rng}."


def excel_layout(path: str, sheet: str, options: str) -> str:
    p = guard(path, ".xlsx")
    wb = load_workbook(str(p))
    ws = _xl_sheet(wb, sheet)
    o = _kv(options)
    done = []
    if "freeze" in o:
        ws.freeze_panes = None if o["freeze"].lower() == "none" else o["freeze"].upper()
        done.append("freeze")
    if "filter" in o:
        if o["filter"].lower() == "none":
            ws.auto_filter.ref = None
        else:
            ws.auto_filter.ref = ws.dimensions if o["filter"].lower() == "auto" else o["filter"].upper()
        done.append("filter")
    if _bool(o.get("auto_width", "")):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[letter].width = min(max(longest + 2, 8), 60)
        done.append("auto_width")
    if "widths" in o:
        for part in o["widths"].split(","):
            col, _, w = part.partition(":")
            if col.strip() and w.strip():
                ws.column_dimensions[col.strip().upper()].width = float(w)
        done.append("widths")
    if "row_height" in o:
        for part in o["row_height"].split(","):
            row, _, h = part.partition(":")
            if row.strip().isdigit() and h.strip():
                ws.row_dimensions[int(row)].height = float(h)
        done.append("row_height")
    if "tab_color" in o:
        ws.sheet_properties.tabColor = _hex(o["tab_color"])
        done.append("tab_color")
    if "gridlines" in o:
        ws.sheet_view.showGridLines = o["gridlines"].lower() not in ("off", "false", "no", "0")
        done.append("gridlines")
    if "zoom" in o:
        ws.sheet_view.zoomScale = int(float(o["zoom"]))
        done.append("zoom")
    if "merge" in o:
        for rng in o["merge"].split(","):
            ws.merge_cells(rng.strip().upper())
        done.append("merge")
    if "unmerge" in o:
        for rng in o["unmerge"].split(","):
            ws.unmerge_cells(rng.strip().upper())
        done.append("unmerge")
    wb.save(str(p))
    return "Layout updated: " + (", ".join(done) or "nothing (no valid options)")



def excel_design_report(path: str, output: str = "", sheet: str = "",
                        title: str = "", subtitle: str = "") -> str:
    """One-call premium redesign of a data sheet. The original file is never modified."""
    src = safe(path)
    if src.suffix.lower() != ".xlsx" or not src.exists():
        raise ValueError(f"File not found: {path}")
    title = title.strip() or src.stem.replace("_", " ").upper()
    subtitle = subtitle.strip() or "DATA REPORT"
    dst = safe(output) if output else src.with_name(src.stem + "_Designed.xlsx")
    if dst.suffix.lower() != ".xlsx":
        raise ValueError("output must end with .xlsx")
    if not confirm(f"Create designed copy: {dst.name} (original stays unchanged)"):
        return "User denied."
    res = design_report(src, dst, sheet, title, subtitle)
    approved.add(str(dst))
    rel = dst.relative_to(ROOT.resolve()).as_posix()
    status = "all checks passed" if res["ok"] else "SOME CHECKS FAILED"
    return f"Saved {rel} ({status}). " + "; ".join(res["checks"])

# ---------- Registry ----------
REGISTRY = [
    ("excel", tool("excel_start", "Create an Excel workbook (optionally copying a template)",
                   {"path": "e.g. reports/Budget.xlsx", "template": "Optional .xlsx template"},
                   ["path"]), excel_start),
    ("excel", tool("excel_add_sheet", "Add a formatted table to a sheet. First line = headers.",
                   {"path": "The .xlsx path", "sheet": "Sheet name",
                    "table": "Header line then rows, cells separated by |",
                    "title": "Optional title above the table"},
                   ["path", "sheet", "table"]), excel_add_sheet),
    ("excel", tool("excel_set_cells", "Set cells or formulas, one per line like B7==SUM(B2:B6) or A7=Total",
                   {"path": "The .xlsx path", "sheet": "Sheet name", "cells": "Lines of REF=VALUE"},
                   ["path", "sheet", "cells"]), excel_set_cells),
    ("excel", tool("excel_add_chart", "Add a native Excel chart",
                   {"path": "The .xlsx path", "sheet": "", "chart_type": "bar, line or pie", "title": "",
                    "data_range": "Values incl. header, e.g. B1:B6", "categories_range": "e.g. A2:A6",
                    "anchor": "Cell, default H2"},
                   ["path", "sheet", "chart_type", "title", "data_range", "categories_range"]),
     excel_add_chart),
    ("excel", tool("excel_read", "Read cells from an existing workbook (use before formatting to see columns and rows)",
                   {"path": "The .xlsx path", "sheet": "Sheet name (blank = first)",
                    "cell_range": "e.g. A1:H20 (blank = top-left area)"}, ["path"]), excel_read),
    ("excel", tool("excel_format",
                   "Style cells in a range. style is key=value pairs joined by |, keys: bold,italic,underline,size,font,color,fill,align,valign,wrap,border,border_color,number_format. Colors: hex or names. border like thin or thin:bottom. Presets: header,title,total,currency,integer,percent,date.",
                   {"path": "The .xlsx path", "sheet": "Sheet name",
                    "cell_range": "e.g. A1:H1, B2:B31, C:C",
                    "style": "e.g. preset header, or bold=true|fill=FFFF00|number_format=0.0"},
                   ["path", "sheet", "cell_range", "style"]), excel_format),
    ("excel", tool("excel_conditional_format",
                   "Add conditional formatting. rule: color_scale, data_bar, greater_than, less_than, equal, not_equal, at_least, at_most, between (value a|b), text_contains, top_n, bottom_n, duplicates, zebra, formula.",
                   {"path": "The .xlsx path", "sheet": "Sheet name", "cell_range": "e.g. C2:C31",
                    "rule": "Rule name", "value": "Threshold, text, n, a|b or formula",
                    "fill": "Fill color (hex or name)", "font_color": "Font color",
                    "colors": "For color_scale: low|mid|high hex colors; for data_bar one color"},
                   ["path", "sheet", "cell_range", "rule"]), excel_conditional_format),
    ("excel", tool("excel_layout",
                   "Sheet layout. options key=value joined by |: freeze=A2, filter=auto, auto_width=true, widths=A:12,B:30, row_height=1:30, tab_color, gridlines=off, zoom=90, merge=A1:H1",
                   {"path": "The .xlsx path", "sheet": "Sheet name", "options": "key=value|key=value"},
                   ["path", "sheet", "options"]), excel_layout),
    ("excel", tool("excel_design_report",
                   "ONE-CALL premium redesign of a data sheet: navy title band, KPI cards (formulas), styled table, rating color scale, zebra rows, filter, frozen header, no gridlines. Saves a NEW file; original untouched; verifies by reopening. Use this for any full redesign request.",
                   {"path": "Source .xlsx", "output": "New file name, e.g. reports/IMDb_Top_30_Movies_Designed.xlsx",
                    "sheet": "Sheet name", "title": "Title text", "subtitle": "Subtitle text"},
                   ["path", "output"]), excel_design_report),
]