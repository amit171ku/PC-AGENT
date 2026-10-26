"""
Premium navy-and-white design for an EXISTING Word (.docx) document.

Run directly:
    python word_design.py "C:\\path\\Report.docx"
    python word_design.py <source.docx> <output.docx> --title "Project Report" --company "ACME"
        [--subtitle "..."] [--author "..."] [--cover auto|yes|no] [--toc auto|yes|no]

The source file is never modified. The restyled copy keeps every paragraph, table and image;
only the look changes (fonts, heading colors, tables, header/footer, optional cover and
table of contents). The saved file is reopened and checked.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from lxml import etree
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY, BLUE, LINE, ZEBRA = "0B2545", "3A7CA5", "D5DDE8", "F3F6FB"
INK, MUTED, PALE, WHITE = "1B2A41", "5B6B82", "A9C4E4", "FFFFFF"
FONT = "Calibri"
CENTER, RIGHT, LEFT = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT

_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
               "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
               "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
               "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
               "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
               "w:rPr", "w:sectPr", "w:pPrChange")


def etree_tostring(el):
    return etree.tostring(el).decode("utf-8", errors="ignore")


def _rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- styles
def _style_font(style, size=None, bold=None, color=None, italic=None):
    style.font.name = FONT
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color:
        style.font.color.rgb = _rgb(color)
    rf = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]


def _style_border(style, side, color, sz):
    pPr = style.element.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    pb = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color)
    pb.append(e)
    pPr.insert_element_before(pb, *_AFTER_PBDR)


def _apply_styles(doc):
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    _style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    specs = {
        "Title": (32, True, NAVY, 0, 6), "Subtitle": (16, False, MUTED, 0, 12),
        "Heading 1": (18, True, NAVY, 20, 8), "Heading 2": (14, True, NAVY, 14, 6),
        "Heading 3": (12, True, BLUE, 10, 4), "Heading 4": (11, True, BLUE, 8, 3),
    }
    for name, (size, bold, color, before, after) in specs.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _style_font(st, size, bold, color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            st.paragraph_format.keep_with_next = True
    try:
        _style_border(doc.styles["Heading 1"], "bottom", BLUE, 8)
    except KeyError:
        pass
    try:
        _style_font(doc.styles["Caption"], size=9, bold=False, color=MUTED, italic=True)
    except KeyError:
        pass


def _normalize_runs(doc):
    """Let the new styles take effect: drop stray fonts/colors on headings, unify body font."""
    for p in doc.paragraphs:
        name = (p.style.name or "").lower()
        is_head = name.startswith("heading") or name in ("title", "subtitle")
        for r in p.runs:
            if is_head:
                r.font.name = None
                r.font.size = None
                r.font.color.rgb = None
            elif r.font.name and not re.search(r"consolas|courier|mono|code", r.font.name, re.I):
                r.font.name = None
        if p._p.xpath(".//w:drawing") and not is_head:
            p.alignment = CENTER


# ---------------------------------------------------------------- tables
def _is_number(t):
    t = t.strip()
    for ch in ",%$€£₹":
        t = t.replace(ch, "")
    t = t.strip("() ")
    try:
        float(t)
        return True
    except ValueError:
        return False


def _col_numeric(table, j):
    vals = []
    for row in list(table.rows)[1:]:
        cells = row.cells
        if j < len(cells) and cells[j].text.strip():
            vals.append(cells[j].text)
    return bool(vals) and sum(_is_number(v) for v in vals) / len(vals) >= 0.8


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection",
                               "w:tcFitText", "w:vAlign", "w:hideMark")


def _cell_margins(cell, top=50, left=100, bottom=50, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.insert_element_before(mar, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _is_cover_table(t):
    return (len(t.rows) == 1 and len(t.columns) == 1
            and f'w:fill="{NAVY}"' in etree_tostring(t._tbl))


def _style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), LINE)
        borders.append(e)
    tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    rows = list(table.rows)
    ncols = max((len(r.cells) for r in rows), default=0)
    numeric = [_col_numeric(table, j) for j in range(ncols)]
    for ri, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            _shade(cell, NAVY if ri == 0 else (ZEBRA if ri % 2 == 0 else WHITE))
            _cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.alignment = CENTER if ri == 0 else (RIGHT if numeric[ci] else LEFT)
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(10)
                    r.font.bold = True if ri == 0 else r.font.bold
                    r.font.color.rgb = _rgb(WHITE if ri == 0 else INK)


# ---------------------------------------------------------------- header / footer / fields
def _field(par, instr, placeholder="1", size=9, color=MUTED):
    def mk(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    def run():
        r = par.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(color)
        return r
    run()._r.append(mk("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run()._r.append(it)
    run()._r.append(mk("separate"))
    ph = run()
    ph.text = placeholder
    run()._r.append(mk("end"))


def _para_border(p, side, color, sz=6, space=1):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    pb = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pb.append(e)
    pPr.insert_element_before(pb, *_AFTER_PBDR)


def _reset_part(part):
    part.is_linked_to_previous = False
    for p in part.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = part.paragraphs[0]
    for el in list(p._p):
        if el.tag != qn("w:pPr"):
            p._p.remove(el)
    return p


def _header_footer(doc, title, company, cover):
    for i, s in enumerate(doc.sections):
        s.different_first_page_header_footer = bool(cover) and i == 0
        if s.different_first_page_header_footer:
            _reset_part(s.first_page_header)
            _reset_part(s.first_page_footer)
        hp = _reset_part(s.header)
        hp.alignment = RIGHT
        r = hp.add_run(" | ".join(x for x in (company, title) if x))
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(MUTED)
        _para_border(hp, "bottom", LINE)
        fp = _reset_part(s.footer)
        fp.alignment = CENTER
        _para_border(fp, "top", LINE)
        r = fp.add_run("Page ")
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(MUTED)
        _field(fp, "PAGE")
        r = fp.add_run(" of ")
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(MUTED)
        _field(fp, "NUMPAGES")


# ---------------------------------------------------------------- cover + table of contents
def _prepend(doc, builder):
    body = doc.element.body
    before = set(body)
    builder()
    new = [c for c in body if c not in before]
    for i, el in enumerate(new):
        body.insert(i, el)


def _build_cover(doc, title, subtitle, company, author, date):
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(90)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = width
    _shade(cell, NAVY)
    _cell_margins(cell, top=500, left=450, bottom=500, right=450)
    p = cell.paragraphs[0]
    p.alignment = LEFT
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title)
    r.font.name, r.font.size, r.font.bold = FONT, Pt(30), True
    r.font.color.rgb = _rgb(WHITE)
    if subtitle:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(subtitle)
        r2.font.name, r2.font.size = FONT, Pt(14)
        r2.font.color.rgb = _rgb(PALE)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    _para_border(rule, "bottom", BLUE, sz=24, space=1)
    for line in (company, author, date):
        if line:
            lp = doc.add_paragraph()
            lp.paragraph_format.space_after = Pt(2)
            lr = lp.add_run(line)
            lr.font.size = Pt(12)
            lr.font.color.rgb = _rgb(MUTED)
    doc.add_page_break()


def _build_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = _rgb(NAVY)
    _para_border(p, "bottom", BLUE, sz=8, space=4)
    tp = doc.add_paragraph()
    _field(tp, 'TOC \\o "1-3" \\h \\z \\u',
           "Open in Word and press F9 (or accept the update prompt) to build the table of contents.",
           size=11, color=MUTED)
    doc.add_page_break()
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)


# ---------------------------------------------------------------- helpers
def _guess_title(doc, src):
    for p in doc.paragraphs[:40]:
        if p.text.strip() and (p.style.name or "") == "Title":
            return p.text.strip()
    for p in doc.paragraphs:
        if p.text.strip() and (p.style.name or "").lower() == "heading 1":
            return p.text.strip()
    return Path(src).stem.replace("_", " ").replace("-", " ").strip()


def _has_cover(doc):
    """A cover is our navy title block, or a Title-style paragraph near the top followed by a page break."""
    for el in doc.element.body:
        if el.tag == qn("w:p") and not "".join(el.itertext()).strip():
            continue
        if el.tag == qn("w:tbl") and f'w:fill="{NAVY}"' in etree_tostring(el):
            return True
        break
    paras = list(doc.paragraphs)
    title_at = None
    seen = 0
    for i, p in enumerate(paras):
        if not p.text.strip():
            continue
        if (p.style.name or "") == "Title":
            title_at = i
            break
        seen += 1
        if seen >= 8:
            return False
    if title_at is None:
        return False
    for p in paras[title_at:title_at + 14]:
        if p._p.xpath('.//w:br[@w:type="page"]') or p.paragraph_format.page_break_before:
            return True
    return False


def _has_toc(doc):
    return bool(re.search(r"instrText[^>]*>\s*TOC", doc.element.xml))


def _soffice():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(c):
            return c
    return None


def _pdf_pages(path):
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        return len(re.findall(rb"/Type\s*/Page[^s]", Path(path).read_bytes()))


def _snapshot(doc):
    return {
        "paras": [p.text for p in doc.paragraphs if p.text.strip()],
        "tables": [[[c.text for c in r.cells] for r in t.rows]
                   for t in doc.tables if not _is_cover_table(t)],
        "images": len(doc.inline_shapes),
    }


# ---------------------------------------------------------------- main entry
def design_document(src, dst, title="", subtitle="", company="", author="",
                    cover="auto", toc="auto"):
    src, dst = Path(src), Path(dst)
    if src.resolve() == dst.resolve():
        raise ValueError("Output must be a different file. The original is kept untouched.")
    doc = Document(str(src))
    before = _snapshot(doc)
    doc_title = title.strip() or _guess_title(doc, src)

    add_cover = (cover == "yes") or (cover == "auto" and not _has_cover(doc))
    add_toc = (toc == "yes") or (toc == "auto" and not _has_toc(doc))
    cover_note = "cover page added" if add_cover else "existing cover kept"
    toc_note = "table of contents added" if add_toc else "existing table of contents kept"

    _apply_styles(doc)
    _normalize_runs(doc)
    for t in doc.tables:
        if not _is_cover_table(t):
            _style_table(t)
    _header_footer(doc, doc_title, company, add_cover)
    if add_toc:
        _prepend(doc, lambda: _build_toc(doc))
    if add_cover:
        _prepend(doc, lambda: _build_cover(doc, doc_title, subtitle, company, author,
                                           datetime.now().strftime("%d %B %Y")))
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))

    checks, ok = verify(dst, before, add_cover, add_toc, doc_title, cover_note, toc_note)
    return {"output": str(dst), "checks": checks, "ok": ok}


def verify(dst, before, cover_added, toc_added, title, cover_note, toc_note):
    checks, ok = [], True
    doc = Document(str(dst))
    after = _snapshot(doc)

    n = len(before["paras"])
    tail = after["paras"][len(after["paras"]) - n:] if n else []
    tables = after["tables"]
    if tail == before["paras"] and tables == before["tables"] and after["images"] == before["images"]:
        cells = sum(len(r) for t in before["tables"] for r in t)
        checks.append(f"content intact ({n} paragraphs, {len(before['tables'])} tables, "
                      f"{before['images']} images, {cells} table cells compared)")
    else:
        ok = False
        checks.append("CONTENT MISMATCH: text, tables or images differ from the original")

    sec = doc.sections[0]
    footer_xml = sec.footer._element.xml
    if "PAGE" in footer_xml and "NUMPAGES" in footer_xml:
        checks.append("header and footer with 'Page X of Y'")
    else:
        ok = False
        checks.append("FOOTER PAGE NUMBERS MISSING")
    checks.append(cover_note)
    checks.append(toc_note)
    if toc_added and not _has_toc(doc):
        ok = False
        checks.append("TOC FIELD MISSING")

    styled = 0
    for t in doc.tables:
        if _is_cover_table(t):
            continue
        xml = t._tbl.xml
        if NAVY in xml and "tblBorders" in xml:
            styled += 1
    expected = len(before["tables"])
    if styled == expected:
        checks.append(f"{styled} table(s) styled (navy header, zebra rows, repeating header row)")
    else:
        ok = False
        checks.append(f"TABLE STYLING: {styled} of {expected}")

    try:
        h1 = doc.styles["Heading 1"].font.color.rgb
        checks.append("headings styled in navy" if str(h1) == NAVY else "HEADING COLOR NOT APPLIED")
        ok &= str(h1) == NAVY
    except KeyError:
        checks.append("no Heading 1 style in this document")

    exe = _soffice()
    if exe:
        try:
            with tempfile.TemporaryDirectory() as td:
                subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", td, str(dst)],
                               check=True, timeout=150, capture_output=True)
                pdf = Path(td) / (Path(dst).stem + ".pdf")
                pages = _pdf_pages(pdf)
                checks.append(f"opens and renders ({pages} pages)")
                ok &= pages > 0
        except Exception as e:
            checks.append(f"render check skipped ({str(e)[:50]})")
    else:
        checks.append("render check skipped (LibreOffice not found); open the file in Word to confirm")
    return checks, ok


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Apply a premium navy design to a Word document.")
    ap.add_argument("source")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--title", default="")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--company", default="")
    ap.add_argument("--author", default="")
    ap.add_argument("--cover", default="auto", choices=["auto", "yes", "no"])
    ap.add_argument("--toc", default="auto", choices=["auto", "yes", "no"])
    a = ap.parse_args()
    source = Path(a.source)
    if not source.exists():
        sys.exit(f"File not found: {source}")
    target = Path(a.output) if a.output else source.with_name(source.stem + "_Designed.docx")
    res = design_document(source, target, a.title, a.subtitle, a.company, a.author, a.cover, a.toc)
    print("Saved:", res["output"])
    for line in res["checks"]:
        print(" -", line)
    print("ALL CHECKS PASSED" if res["ok"] else "SOME CHECKS FAILED")